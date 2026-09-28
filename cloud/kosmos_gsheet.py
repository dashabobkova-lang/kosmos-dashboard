# -*- coding: utf-8 -*-
"""Заливает собранный xlsx в Google Таблицу КОСМОС через сервисный аккаунт.

  python kosmos_gsheet.py <файл.xlsx>

Обновляет только листы отчёта; служебный лист скрипта (_данные) не трогает.
Градиенты на колонках дат пересоздаются под текущую ширину листа.
"""
import datetime
import os
import re
import sys

import gspread
from openpyxl import load_workbook

KEY_PATH = os.environ.get("GSHEET_KEY_PATH", "service_account.json")
SHEET_ID = os.environ.get("KOSMOS_SHEET_ID", "1mLPTkAl84cYwKKPZN_IsnrSuMPYA3vXweajwoup5e80")
SKIP_SHEETS = {"_данные"}          # кэш Apps Script

DATE_LIKE = re.compile(r"^\d{1,2}[-.]\d{1,2}$")
GRADIENT_MIN = {"red": 1.0, "green": 0.953, "blue": 0.902}
GRADIENT_MAX = {"red": 0.961, "green": 0.486, "blue": 0.0}
HEADER_DATE_ROW = 2
DATA_START_ROW = 5


def to_ru_formula(f):
    """В русской локали разделитель аргументов — точка с запятой."""
    out, in_str = [], False
    for ch in f:
        if ch == '"':
            in_str = not in_str
        out.append(";" if ch == "," and not in_str else ch)
    return "".join(out)


def cell_value(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        # шапка дней хранится датой с форматом dd.mm - отдаём текстом
        return "'" + v.strftime("%d.%m")
    if isinstance(v, str):
        if v.startswith("="):
            return to_ru_formula(v)
        if DATE_LIKE.match(v):
            return "'" + v
    return v


def date_columns(rows):
    if len(rows) < HEADER_DATE_ROW:
        return []
    head = rows[HEADER_DATE_ROW - 1]
    return [i for i, v in enumerate(head)
            if DATE_LIKE.match(str(v or "").lstrip("'").strip())]


def group_runs(cols):
    runs, cur = [], []
    for c in cols:
        if cur and c == cur[-1] + 1:
            cur.append(c)
        else:
            if cur:
                runs.append(cur)
            cur = [c]
    if cur:
        runs.append(cur)
    return runs


def refresh_gradients(sh, ws, rows):
    cols = date_columns(rows)
    if not cols or len(rows) < DATA_START_ROW:
        return 0
    meta = sh.fetch_sheet_metadata()
    existing = 0
    for sheet in meta.get("sheets", []):
        if sheet.get("properties", {}).get("sheetId") == ws.id:
            existing = len(sheet.get("conditionalFormats", []) or [])
            break
    requests = [{"deleteConditionalFormatRule": {"index": 0, "sheetId": ws.id}}
                for _ in range(existing)]
    for run in group_runs(cols):
        requests.append({"addConditionalFormatRule": {"index": 0, "rule": {
            "ranges": [{"sheetId": ws.id,
                        "startRowIndex": DATA_START_ROW - 1, "endRowIndex": len(rows),
                        "startColumnIndex": run[0], "endColumnIndex": run[-1] + 1}],
            "gradientRule": {"minpoint": {"color": GRADIENT_MIN, "type": "MIN"},
                             "maxpoint": {"color": GRADIENT_MAX, "type": "MAX"}}}}})
    if requests:
        sh.batch_update({"requests": requests})
    return len(group_runs(cols))


PINK = {"red": 0.761, "green": 0.196, "blue": 0.478}        # C2327A
PINK_LIGHT = {"red": 0.949, "green": 0.863, "blue": 0.914}  # F2DCE9
WHITE = {"red": 1, "green": 1, "blue": 1}
BLACK = {"red": 0, "green": 0, "blue": 0}


GREY = {"red": 0.925, "green": 0.925, "blue": 0.925}


def group_key(art):
    a = str(art or "").strip()
    return a.split("/")[0].strip() if "/" in a else a


def body_format(sh, ws, rows):
    """Тело таблицы: числовой формат на всю ширину и чередование групп артикулов.

    Красим по текущей ширине, поэтому новые колонки дней сразу получают
    тот же вид, что и старые.
    """
    width = max(len(r) for r in rows)
    n = len(rows)
    if n < DATA_START_ROW:
        return
    requests = [
        # строки ИТОГО, ₽ и ИТОГО, шт
        {"repeatCell": {
            "range": {"sheetId": ws.id, "startRowIndex": 2, "endRowIndex": 4,
                      "startColumnIndex": 0, "endColumnIndex": width},
            "cell": {"userEnteredFormat": {
                "backgroundColor": GREY,
                "textFormat": {"bold": True},
                "numberFormat": {"type": "NUMBER", "pattern": "# ##0"}}},
            "fields": "userEnteredFormat(backgroundColor,textFormat,numberFormat)"}},
        # данные: числовой формат
        {"repeatCell": {
            "range": {"sheetId": ws.id, "startRowIndex": DATA_START_ROW - 1, "endRowIndex": n,
                      "startColumnIndex": 5, "endColumnIndex": width},
            "cell": {"userEnteredFormat": {
                "numberFormat": {"type": "NUMBER", "pattern": "# ##0"}}},
            "fields": "userEnteredFormat.numberFormat"}},
    ]

    # чередование по группам артикулов (колонка B)
    band, prev, start = False, None, DATA_START_ROW - 1
    blocks = []
    for i in range(DATA_START_ROW - 1, n):
        art = rows[i][1] if len(rows[i]) > 1 else ""
        g = group_key(art)
        if prev is None:
            prev = g
        elif g != prev:
            blocks.append((start, i, band))
            band, prev, start = not band, g, i
    blocks.append((start, n, band))
    for a, b, painted in blocks:
        requests.append({"repeatCell": {
            "range": {"sheetId": ws.id, "startRowIndex": a, "endRowIndex": b,
                      "startColumnIndex": 0, "endColumnIndex": width},
            "cell": {"userEnteredFormat": {"backgroundColor": GREY if painted else WHITE}},
            "fields": "userEnteredFormat.backgroundColor"}})

    for i in range(0, len(requests), 150):
        sh.batch_update({"requests": requests[i:i + 150]})


def head_format(sh, ws, rows):
    """Шапка: дни - бледно-розовые, служебные колонки - малиновые.

    Google не сдвигает старое форматирование при добавлении дня, поэтому
    красим строку заголовков заново на всю текущую ширину.
    """
    if len(rows) < HEADER_DATE_ROW:
        return
    days = set(date_columns(rows))
    width = max(len(r) for r in rows)
    requests = []
    for col in range(width):
        is_day = col in days
        requests.append({"repeatCell": {
            "range": {"sheetId": ws.id,
                      "startRowIndex": HEADER_DATE_ROW - 1, "endRowIndex": HEADER_DATE_ROW,
                      "startColumnIndex": col, "endColumnIndex": col + 1},
            "cell": {"userEnteredFormat": {
                "backgroundColor": PINK_LIGHT if is_day else PINK,
                "horizontalAlignment": "CENTER",
                "verticalAlignment": "MIDDLE",
                "wrapStrategy": "WRAP",
                "textFormat": {"bold": True, "foregroundColor": BLACK if is_day else WHITE}}},
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,"
                      "verticalAlignment,wrapStrategy,textFormat)"}})
    # склейка соседних одинаковых колонок не нужна - запросов немного
    for i in range(0, len(requests), 200):
        sh.batch_update({"requests": requests[i:i + 200]})


def main():
    path = sys.argv[1]
    gc = gspread.service_account(filename=KEY_PATH)
    sh = gc.open_by_key(SHEET_ID)
    wb = load_workbook(path)
    existing = {ws.title: ws for ws in sh.worksheets()}

    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        src = wb[name]
        rows = [[cell_value(c) for c in row] for row in src.iter_rows()]
        n_rows, n_cols = src.max_row, src.max_column

        ws = existing.get(name)
        if ws is None:
            ws = sh.add_worksheet(title=name, rows=n_rows + 10, cols=n_cols + 5)
        else:
            if ws.row_count < n_rows or ws.col_count < n_cols:
                ws.resize(rows=max(ws.row_count, n_rows + 10),
                          cols=max(ws.col_count, n_cols + 5))
            ws.batch_clear([f"A1:{gspread.utils.rowcol_to_a1(ws.row_count, ws.col_count)}"])

        ws.update(range_name="A1", values=rows, value_input_option="USER_ENTERED")
        if name != "Запас":
            body_format(sh, ws, rows)     # сначала фон, потом градиенты поверх
        runs = refresh_gradients(sh, ws, rows)
        head_format(sh, ws, rows)
        print("%-16s %dx%d, градиентов %d" % (name, n_rows, n_cols, runs))

    print("готово -> https://docs.google.com/spreadsheets/d/%s/edit" % SHEET_ID)


if __name__ == "__main__":
    main()
